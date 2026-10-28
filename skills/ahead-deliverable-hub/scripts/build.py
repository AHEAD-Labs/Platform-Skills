#!/usr/bin/env python3
"""Build a self-contained AHEAD engagement hub (one HTML file) from a JSON config.

Examples
  python3 build.py config.json --out "Client x AHEAD Phase 2 Deliverable Package v1.html" --guide "How to add content.md"
  python3 build.py --preset target-design --client "Contoso" --phase-number 2 --date "6 November 2026" --out hub.html
  python3 build.py --preset discovery --dump-config my_config.json      # copy a preset to edit

The config format is described in references/config-schema.md. Any string may use the tokens
{client} {engagement} {phase} {next_phase} {prev_phase} {p} {np} {date} {package}.
"""
import argparse, base64, copy, csv, datetime as dt, json, mimetypes, pathlib, re, sys
from hublib import (e, md, mdi, ph, tag, initials, make_tokens, subst, slot_open, slot_close,
                    HOW_TEXT, HOW_TABLE, HOW_STATS, HOW_DATA, ICONS)

ROOT = pathlib.Path(__file__).resolve().parent.parent
ASSETS = ROOT / 'assets'
SKILL_NAME = 'ahead-deliverable-hub'
VERSION = '1.3.5'


def look_hash():
    """Fingerprint of the bundled styles and behavior, stamped into every hub and verified by check.py."""
    import hashlib
    return hashlib.sha256(((ASSETS / 'hub.css').read_text() + (ASSETS / 'hub.js').read_text()).encode()).hexdigest()[:12]
PALETTE = ['#009FDC', '#00B388', '#7DC8E8', '#FFC47A', '#5FD0B4', '#BFE3F2', '#93AAC0']
ARROW = '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M5 12h14M13 6l6 6-6 6"/></svg>'
KIND_LABELS = {'hypothesis': ('Hypothesis', 'Hypotheses'), 'decision': ('Decision', 'Decisions'),
               'experiment': ('Experiment', 'Experiments'), 'risk': ('Risk', 'Risks'),
               'assumption': ('Assumption', 'Assumptions'), 'question': ('Open question', 'Open questions'),
               'finding': ('Finding', 'Findings'), 'requirement': ('Requirement', 'Requirements')}
FACET_LABELS = {'kind': 'Type', 'priority': 'Priority', 'status': 'Status', 'area': 'Area', 'confidence': 'Confidence', 'owner': 'Owner'}


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', str(s).lower()).strip('-') or 'x'


class Hub:
    def __init__(self, cfg, base):
        self.base = base
        self.tokens = make_tokens(cfg)
        self.cfg = subst(cfg, self.tokens)
        self.owners = self.cfg.get('owners', {})
        self.default_owner = self.cfg.get('default_owner', 'Owner to confirm')
        self.keys = set()
        self.auto = 0
        self.tabs = [self.expand(t) for t in self.cfg['tabs']]
        self.tabmap = {t['key']: t for t in self.tabs}

    # ---------------------------------------------------------- helpers
    def who(self, o):
        o = o or self.default_owner
        return self.owners.get(o, o)

    def key(self, b, tab, sec, suffix=''):
        k = b.get('key')
        if not k:
            self.auto += 1
            k = f"{tab['key']}-{sec.get('id', 's')}{suffix or '-' + str(self.auto)}"
        if k in self.keys:
            raise SystemExit(f'Duplicate slot key "{k}". Every slot key must be unique.')
        self.keys.add(k)
        return k

    def label(self, b, tab, sec):
        return b.get('label') or f"{tab.get('label')} > {sec.get('title', '')}".strip(' >')

    def datauri(self, path):
        p = (self.base / path).resolve() if not pathlib.Path(path).is_absolute() else pathlib.Path(path)
        if not p.exists():
            raise SystemExit(f'Image not found: {p}')
        mime = mimetypes.guess_type(str(p))[0] or 'image/png'
        return f'data:{mime};base64,' + base64.b64encode(p.read_bytes()).decode()

    # ---------------------------------------------------------- deliverable shorthand
    def expand(self, tab):
        tab = copy.deepcopy(tab)
        d = tab.get('deliverable')
        if not d:
            return tab
        k = tab['key']
        own = d.get('owner')
        tab.setdefault('id', d.get('id'))
        tab.setdefault('title', d.get('title', tab.get('label')))
        tab.setdefault('question', d.get('question', ''))
        tab.setdefault('owner', own)
        tab['hero'] = {'chip': d.get('id'), 'title': d.get('title', tab.get('label')), 'lede': d.get('question', '')}
        secs = list(tab.get('sections_before', []))
        if d.get('summary', True):
            secs.append({'id': 'summary', 'eyebrow': 'Summary', 'title': d.get('summary_title', 'Key conclusions'),
                         'blocks': [{'type': 'slot', 'key': f'{k}-summary-text', 'owner': own, 'content': d.get('summary_content'),
                                     'guide': d.get('summary_guide', 'A concise answer to the deliverable question above, in a few sentences.')}]})
        for s in d.get('sections', []):
            if isinstance(s, (list, tuple)):
                s = dict(zip(('key', 'title', 'guide'), s))
            sk = s['key']
            so = s.get('owner', own)
            base = {'key': f'{k}-{sk}-body', 'owner': so, 'guide': s.get('guide', ''), 'content': s.get('content')}
            if s.get('columns'):
                blk = dict(base, type='table', columns=s['columns'], rows=s.get('rows'))
            elif s.get('kind') == 'diagram':
                blk = dict(base, type='figure', src=s.get('src'), caption=s.get('caption'), alt=s.get('alt'))
            else:
                blk = dict(base, type='slot', kind=s.get('kind', 'text'))
            secs.append({'id': sk, 'eyebrow': s.get('eyebrow'), 'title': s['title'], 'intro': s.get('intro'),
                         'blocks': [blk] + s.get('blocks', [])})
        if d.get('decisions', True):
            cols = d.get('decision_columns', ['ID', 'Decision or open item', 'Outcome or next step', '{client} owner', 'Date'])
            cols = [subst(c, self.tokens) for c in cols]
            secs.append({'id': 'decisions', 'eyebrow': 'Decisions', 'title': 'Decisions and open items',
                         'blocks': [{'type': 'table', 'key': f'{k}-decisions-table', 'owner': own, 'columns': cols,
                                     'rows': d.get('decision_rows'),
                                     'guide': 'Add one row per decision or open item for this deliverable.'}]})
        if d.get('criteria'):
            secs.append({'id': 'acceptance', 'eyebrow': 'Acceptance', 'title': 'Acceptance criteria coverage',
                         'intro': 'Each acceptance criterion for this deliverable, mapped to the section that addresses it.',
                         'blocks': [{'type': 'acceptance', 'criteria': d['criteria']}]})
        tab['sections'] = secs + list(tab.get('sections_after', []))
        return tab

    # ---------------------------------------------------------- blocks
    def block(self, b, tab, sec):
        fn = getattr(self, 'b_' + b.get('type', 'slot').replace('-', '_'), None)
        if not fn:
            raise SystemExit(f"Unknown block type '{b.get('type')}' in tab '{tab['key']}'")
        return fn(b, tab, sec)

    def b_slot(self, b, tab, sec):
        k = self.key(b, tab, sec)
        o = self.who(b.get('owner', tab.get('owner')))
        inner = md(b['content']) if b.get('content') else ph(b.get('guide', ''), o, b.get('kind', 'text'))
        return slot_open(k, self.label(b, tab, sec), o, HOW_TEXT) + inner + slot_close(k)

    def b_text(self, b, tab, sec):
        return f'<div class="prose">{md(b.get("content", ""))}</div>'

    def b_callout(self, b, tab, sec):
        return f'<div class="callout">{mdi(b.get("content", ""))}</div>'

    def b_figure(self, b, tab, sec):
        k = self.key(b, tab, sec)
        o = self.who(b.get('owner', tab.get('owner')))
        if b.get('src') or b.get('svg'):
            img = b['svg'] if b.get('svg') else f'<img src="{self.datauri(b["src"])}" alt="{e(b.get("alt", b.get("caption", "")))}">'
            cap = f'<figcaption>{mdi(b["caption"])}</figcaption>' if b.get('caption') else ''
            inner = f'<figure>{img}{cap}</figure>'
        elif b.get('content'):
            inner = md(b['content'])
        else:
            inner = ph(b.get('guide', ''), o, 'diagram')
        return slot_open(k, self.label(b, tab, sec), o, HOW_TEXT) + inner + slot_close(k)

    def rows_html(self, rows, strong_first=False):
        out = ''
        for r in rows or []:
            cells = r if isinstance(r, (list, tuple)) else [r]
            out += '<tr>' + ''.join(
                f'<td>{"<b>" if (strong_first and i == 0) else ""}{mdi(c)}{"</b>" if (strong_first and i == 0) else ""}</td>'
                for i, c in enumerate(cells)) + '</tr>\n'
        return out

    def b_table(self, b, tab, sec):
        cols = b.get('columns', [])
        head = '<thead><tr>' + ''.join(f'<th>{e(c)}</th>' for c in cols) + '</tr></thead>'
        rows = self.rows_html(b.get('rows'), b.get('strong_first', False))
        if not b.get('key') and rows and not b.get('editable'):
            return f'<div class="tw"><table>{head}<tbody>\n{rows}</tbody></table></div>'
        k = self.key(b, tab, sec)
        o = self.who(b.get('owner', tab.get('owner')))
        if not rows:
            rows = f'  <tr class="ph-row"><td colspan="{max(len(cols), 1)}">{ph(b.get("guide", ""), o, "table")}</td></tr>\n'
        return (slot_open(k, self.label(b, tab, sec), o, HOW_TABLE)
                + f'<div class="tw"><table>{head}<tbody>\n{rows}</tbody></table></div>' + slot_close(k))

    def b_stats(self, b, tab, sec):
        k = self.key(b, tab, sec, '-stats')
        o = self.who(b.get('owner', tab.get('owner')))
        items = b.get('items', [])
        cards = ''
        for it in items:
            v = it.get('value')
            n = e(v) if v not in (None, '') else f'<span class="skel" data-owner="{e(o)}" title="Number to add"></span>'
            cards += f'<div class="stat"><div class="n">{n}</div><div class="l">{e(it.get("label", ""))}</div><div class="s">{e(it.get("sub", ""))}</div></div>'
        return (slot_open(k, self.label(b, tab, sec) if b.get('label') else 'Headline numbers', o, HOW_STATS)
                + f'<div class="stats" style="--n:{min(max(len(items), 1), 4)}">{cards}</div>' + slot_close(k))

    def b_cards(self, b, tab, sec):
        cols = b.get('cols', 2)
        out = f'<div class="grid{3 if cols == 3 else 2} es">'
        for i, c in enumerate(b.get('items', [])):
            ac = c.get('accent', PALETTE[i % len(PALETTE)])
            eb = f'<p class="eyebrow">{e(c["eyebrow"])}</p>' if c.get('eyebrow') else ''
            inner = dict(c, type=c.get('type', 'slot'), label=c.get('label') or f"{tab.get('label')} > {sec.get('title', '')} > {c.get('title', '')}")
            out += f'<div class="card accent" style="--ac:{ac}">{eb}<h3>{e(c.get("title", ""))}</h3>{self.block(inner, tab, sec)}</div>'
        return out + '</div>'

    def b_lead(self, b, tab, sec):
        eb = f'<p class="eyebrow">{e(b["eyebrow"])}</p>' if b.get('eyebrow') else ''
        h = f'<h3>{e(b["title"])}</h3>' if b.get('title') else ''
        lab = b.get('label') or ' > '.join(x for x in (tab.get('label'), sec.get('title'), b.get('eyebrow') or b.get('title')) if x)
        inner = self.block(dict(b, type=b.get('inner', 'slot'), label=lab), tab, sec)
        note = f'<p class="reveal">{mdi(b["note"])}</p>' if b.get('note') else ''
        return f'<div class="lead-card">{eb}{h}{inner}{note}</div>'

    def b_index(self, b, tab, sec):
        keys = b.get('tabs') or [t['key'] for t in self.tabs if t['key'] != tab['key'] and t.get('index', True)]
        out = '<div class="dx">'
        for i, k in enumerate(keys):
            t = self.tabmap[k]
            ac = t.get('accent', PALETTE[i % len(PALETTE)])
            tid = t.get('id')
            chip = f'<span class="dx-id">{e(tid)}</span>' if tid else ''
            q = t.get('question') or t.get('summary') or ''
            out += (f'<button class="dx-row{"" if tid else " noid"}" type="button" onclick="showTab(\'{k}\')" style="--ac:{ac}">{chip}'
                    f'<span class="dx-main"><span class="dx-t">{e(t.get("title", t["label"]))}</span><span class="dx-q">{e(q)}</span></span>'
                    f'<span class="dx-go" aria-hidden="true">{ARROW}</span></button>')
        return out + '</div>'

    def b_glossary(self, b, tab, sec):
        k = self.key(b, tab, sec, '-terms')
        o = self.who(b.get('owner', tab.get('owner')))
        terms = b.get('terms')
        if terms:
            inner = '<dl class="gl">' + ''.join(f'<dt>{e(t)}</dt><dd>{mdi(d)}</dd>' for t, d in terms) + '</dl>'
        else:
            inner = ph(b.get('guide', 'Definitions of the terms and acronyms used across the hub.'), o, 'text')
        return slot_open(k, self.label(b, tab, sec), o, HOW_TEXT) + inner + slot_close(k)

    def b_details(self, b, tab, sec):
        inner = ''.join(self.block(x, tab, sec) for x in b.get('blocks', []))
        return f'<details class="dt"{" open" if b.get("open") else ""}><summary>{e(b.get("summary", "More detail"))}</summary><div class="dt-b">{inner}</div></details>'

    def b_acceptance(self, b, tab, sec):
        titles = {s['id']: s.get('title', s['id']) for s in tab.get('sections', [])}
        rows = ''
        for crit, ks in b.get('criteria', []):
            links = ', '.join(
                f'<a href="#{tab["key"]}-{k}" onclick="return goAnchor(\'{tab["key"]}-{k}\')">{e(titles.get(k, k))}</a>' for k in ks)
            rows += f'<tr><td>{mdi(crit)}</td><td>{links}</td></tr>'
        return f'<div class="tw"><table><thead><tr><th>Acceptance criterion</th><th>Covered in</th></tr></thead><tbody>{rows}</tbody></table></div>'

    # ---------------------------------------------------------- gantt
    def weeks(self, spec):
        if isinstance(spec, list):
            return [tuple(w) for w in spec], len(spec)
        n, after = spec.get('count', 4), spec.get('after', 0)
        try:
            start = dt.date.fromisoformat(str(spec.get('start', '')))
        except ValueError:
            start = None  # no start date given: show week numbers only
        uk = str(self.cfg.get('lang', 'en-US')).lower().startswith('en-gb')
        fmt = lambda d: (f'{d.day} {d.strftime("%b")}' if uk else f'{d.strftime("%b")} {d.day}') if start else ''
        at = lambda i: fmt(start + dt.timedelta(weeks=i)) if start else ''
        ws = [(f'Wk {i + 1}', at(i)) for i in range(n)]
        ws += [(f'+{j + 1} wk', at(n + j)) for j in range(after)]
        return ws, n

    def b_gantt(self, b, tab, sec):
        ws, active = self.weeks(b['weeks'])
        active = b.get('active', active)
        N = len(ws)

        def cells(a, z, marks):
            h = ''
            for i in range(1, N + 1):
                inner = ''
                if a and a <= i <= z:
                    cls = '' if a == z else ('l' if i == a else ('r' if i == z else 'm'))
                    inner += f'<div class="g-bar {cls}"></div>'
                for w, t in marks:
                    if w == i:
                        inner += {'due': '<div class="g-due" title="Deliverable due"></div>',
                                  'ms': '<div class="g-ms" title="Milestone"></div>',
                                  'star': '<span class="g-star" title="Decision point">&#9733;</span>'}[t]
                h += f'<div class="g-cell{" post" if i > active else ""}">{inner}</div>'
            return h
        key = ('<div class="g-key"><span><i class="k-bar"></i>Task</span><span><i class="k-due"></i>Deliverable due</span>'
               '<span><span style="color:#FFC47A">&#9733;</span>Decision point</span><span><i class="k-ms"></i>Milestone</span></div>')
        h = key + f'<div class="gantt" style="--wk:{N}"><div class="g-row head"><div>{e(b.get("first_column", "Deliverable and key tasks"))}</div><div>Date</div>' + ''.join(
            f'<div class="wk">{e(a)}<span>{e(z)}</span></div>' for a, z in ws) + '</div>'
        for r in b.get('rows', []):
            due = [(r['due'], 'due')] if r.get('due') else []
            h += f'<div class="g-row del"><div class="g-name"><span class="id">{e(r.get("id", ""))}</span>{e(r["title"])}</div><div class="g-tgt">{e(r.get("date", ""))}</div>{cells(0, 0, due)}</div>'
            for t in r.get('tasks', []):
                a, z = t['from'], t.get('to', t['from'])
                when = t.get('date') or (f'Week {a}' if a == z else f'Weeks {a} to {z}')
                marks = [(t['decision'], 'star')] if t.get('decision') else []
                h += f'<div class="g-row"><div class="g-name"><span class="id">{e(t.get("id", ""))}</span>{e(t["title"])}</div><div class="g-tgt">{e(when)}</div>{cells(a, z, marks)}</div>'
        if b.get('milestones'):
            h += '<div class="g-row grp"><div>Milestones</div></div>'
            for m in b['milestones']:
                h += f'<div class="g-row"><div class="g-name"><span class="id">{e(m.get("id", ""))}</span>{e(m["title"])}</div><div class="g-tgt">{e(m.get("date", ""))}</div>{cells(0, 0, [(m["week"], "ms")])}</div>'
        return h + '</div>'

    # ---------------------------------------------------------- register
    def load_items(self, path):
        p = self.base / path
        rows = []
        if p.suffix.lower() in ('.xlsx', '.xlsm'):
            import openpyxl
            ws = openpyxl.load_workbook(p, data_only=True).active
            data = [[c for c in r] for r in ws.iter_rows(values_only=True)]
            hi = next(i for i, r in enumerate(data) if any(str(c or '').strip().lower() == 'id' for c in r))
            hdr = [str(c or '').strip() for c in data[hi]]
            rows = [dict(zip(hdr, r)) for r in data[hi + 1:] if any(c not in (None, '') for c in r)]
        else:
            with open(p, newline='', encoding='utf-8-sig') as f:
                rows = list(csv.DictReader(f))
        alias = {'id': 'id', 'type': 'kind', 'kind': 'kind', 'title': 'title', 'name': 'title', 'statement': 'statement',
                 'description': 'statement', 'hypothesis': 'statement', 'priority': 'priority', 'status': 'status',
                 'area': 'area', 'domain': 'area', 'architecture component': 'area', 'component': 'area', 'owner': 'owner',
                 'evidence': 'evidence', 'source': 'evidence', 'sources': 'evidence', 'implication': 'implication',
                 'next step': 'next', 'next steps': 'next', 'next': 'next', 'confidence': 'confidence', 'outcome': 'outcome'}
        items = []
        for r in rows:
            it = {}
            for k, v in r.items():
                f = alias.get(str(k).strip().lower())
                if f and v not in (None, ''):
                    it[f] = str(v).strip()
            if it.get('id'):
                it.setdefault('kind', {'H': 'hypothesis', 'D': 'decision', 'E': 'experiment', 'R': 'risk', 'A': 'assumption', 'Q': 'question'}.get(it['id'][:1].upper(), 'finding'))
                it['kind'] = it['kind'].lower().rstrip('s') if it['kind'].lower() not in KIND_LABELS else it['kind'].lower()
                items.append(it)
        return items

    def b_register(self, b, tab, sec):
        k = self.key(b, tab, sec, '-register')
        o = self.who(b.get('owner', tab.get('owner')))
        items = b.get('items') or (self.load_items(b['items_file']) if b.get('items_file') else [])
        if not items:
            return slot_open(k, self.label(b, tab, sec), o, HOW_DATA) + ph(b.get('guide', ''), o, 'data') + slot_close(k)
        facets = [f for f in b.get('facets', ['kind', 'priority', 'status', 'area']) if any(it.get(f) for it in items)]
        orders = b.get('orders', {'priority': ['High', 'Medium', 'Low']})
        chips = ''
        for f in facets:
            vals = {}
            for it in items:
                if it.get(f):
                    vals[it[f]] = vals.get(it[f], 0) + 1
            order = orders.get(f, [])
            keys = sorted(vals, key=lambda v: (order.index(v) if v in order else 99, v))
            lab = (lambda v: KIND_LABELS.get(v, (v.title(), v.title() + 's'))[1]) if f == 'kind' else (lambda v: v)
            chips += (f'<div class="rg-g"><span class="rg-gl">{e(FACET_LABELS.get(f, f.title()))}</span>'
                      + ''.join(f'<button class="rg-chip" type="button" aria-pressed="false" data-f="{f}" data-v="{e(slug(v))}">{e(lab(v))}<i>{vals[v]}</i></button>' for v in keys) + '</div>')
        def tcls(f, v):
            v2 = str(v).lower()
            if f == 'priority':
                return {'high': 'warn', 'medium': 'info'}.get(v2, '')
            if f == 'status':
                if any(w in v2 for w in ('closed', 'complete', 'agreed', 'validated', 'proven', 'resolved')):
                    return 'ok'
                if any(w in v2 for w in ('defer', 'blocked', 'risk', 'partial')):
                    return 'warn'
                if any(w in v2 for w in ('rejected', 'not proven', 'failed', 'disproven')):
                    return 'bad'
                return 'info'
            return ''
        lst = ''
        for it in items:
            kind = it.get('kind', 'finding')
            tags = ''.join(f'<span class="tag t-{tcls(f, it[f])}">{e(it[f])}</span>' for f in ('priority', 'status') if it.get(f))
            meta = ''
            for f, lab in (('evidence', 'Evidence'), ('implication', 'Implication'), ('outcome', 'Outcome'), ('next', 'Next step'),
                           ('area', 'Area'), ('confidence', 'Confidence'), ('owner', 'Owner')):
                v = it.get(f)
                if v:
                    if isinstance(v, list):
                        v = '<ul>' + ''.join(f'<li>{mdi(x)}</li>' for x in v) + '</ul>'
                    else:
                        v = mdi(v)
                    meta += f'<dt>{lab}</dt><dd>{v}</dd>'
            data = ' '.join(f'data-{f}="{e(slug(it[f]))}"' for f in facets if it.get(f))
            stmt = f'<p class="ri-s">{mdi(it["statement"])}</p>' if it.get('statement') else ''
            klab = KIND_LABELS.get(kind, (kind.title(),))[0]
            lst += (f'<details class="ri" id="ri-{e(it["id"])}" {data}><summary><span class="ri-id">{e(it["id"])}</span>'
                    f'<span class="ri-t">{mdi(it.get("title", ""))}</span><span class="ri-tags"><span class="tag">{e(klab)}</span>{tags}</span></summary>'
                    f'<div class="ri-b">{stmt}<dl class="ri-m">{meta}</dl></div></details>\n')
        body = (f'<div class="rg" data-reg><div class="rg-bar"><input class="rg-q" type="search" placeholder="Search the register" aria-label="Search the register">'
                f'<span class="rg-n"></span></div><div class="rg-f">{chips}</div><div class="rg-list">\n{lst}</div></div>')
        return slot_open(k, self.label(b, tab, sec), o, 'To update: fill_slot.py FILE --key <this key> --block register.json (or rebuild while the file is unedited).') + body + slot_close(k)

    # ---------------------------------------------------------- heatmap
    def b_heatmap(self, b, tab, sec):
        k = self.key(b, tab, sec, '-heatmap')
        o = self.who(b.get('owner', tab.get('owner')))
        scores = b.get('scores')
        if not scores:
            return slot_open(k, self.label(b, tab, sec), o, HOW_DATA) + ph(b.get('guide', ''), o, 'data') + slot_close(k)
        sc = b.get('scale', {})
        lo, hi = sc.get('min', 0), sc.get('max', 4)
        labels = {str(x): y for x, y in sc.get('labels', {}).items()}
        notes = b.get('notes', {})
        rows, cols = b['rows'], b['columns']
        def colour(v):
            t = (float(v) - lo) / ((hi - lo) or 1)
            t = max(0.0, min(1.0, t))
            return f'rgba(0,159,220,{0.10 + 0.72 * t:.2f})'
        head = '<thead><tr><th></th>' + ''.join(f'<th>{e(c)}</th>' for c in cols) + '</tr></thead>'
        body = ''
        for r, row in zip(rows, scores):
            body += f'<tr><th class="hm-r" scope="row">{e(r)}</th>'
            for c, v in zip(cols, row):
                n = notes.get(f'{r}|{c}', '')
                if v in (None, ''):
                    body += f'<td class="hm-c hm-na" title="{e(r)} · {e(c)}: not assessed{(". " + e(n)) if n else ""}"></td>'
                else:
                    lab = labels.get(str(v).rstrip('0').rstrip('.') if isinstance(v, float) else str(v), '')
                    tip = f'{r} · {c}: {v}{(" (" + lab + ")") if lab else ""}{(". " + n) if n else ""}'
                    body += f'<td class="hm-c" style="background:{colour(v)}" title="{e(tip)}">{e(v)}</td>'
            body += '</tr>'
        legend = ''
        if labels:
            legend = '<div class="hm-legend">' + ''.join(
                f'<span><i style="background:{colour(x)}"></i>{e(x)} {e(y)}</span>' for x, y in labels.items()) + '<span><i class="hm-na" style="background:repeating-linear-gradient(135deg,rgba(125,200,232,.18) 0 4px,transparent 4px 8px)"></i>Not assessed</span></div>'
        return (slot_open(k, self.label(b, tab, sec), o, 'To update: fill_slot.py FILE --key <this key> --block heatmap.json (or rebuild while the file is unedited).')
                + f'<div class="tw"><table class="hm">{head}<tbody>{body}</tbody></table></div>{legend}' + slot_close(k))

    # ---------------------------------------------------------- page structure
    def section(self, s, tab):
        sid = s['id'] if s['id'].startswith(tab['key'] + '-') else f"{tab['key']}-{s['id']}"
        if sid in self.keys:
            raise SystemExit(f'Section id "{sid}" clashes with another section or slot key. Rename one of them.')
        self.keys.add(sid)
        eb = f'<p class="eyebrow">{e(s["eyebrow"])}</p>' if s.get('eyebrow') else ''
        intro = f'<p>{mdi(s["intro"])}</p>' if s.get('intro') else ''
        inner = ''.join(self.block(b, tab, s) for b in s.get('blocks', []))
        return f'<section class="sec" id="{e(sid)}"><div class="wrap"><div class="sh">{eb}<h2>{e(s["title"])}</h2>{intro}</div>{inner}</div></section>\n'

    def hero(self, tab):
        c = tab.get('cover')
        if c:
            stats = self.block(dict(c['stats'], type='stats'), tab, {'id': 'cover', 'title': 'Headline numbers'}) if c.get('stats') else ''
            return (f'<section class="hero cover"><div class="wrap"><p class="eyebrow">{e(c.get("eyebrow", ""))}</p>'
                    f'<h1>{e(c.get("title", ""))}</h1><p class="lede">{mdi(c.get("lede", ""))}</p>{stats}</div></section>\n')
        h = tab.get('hero') or {'title': tab.get('title', tab['label']), 'lede': tab.get('question', '')}
        top = (f'<p style="margin:0 0 18px"><span class="chip-id">{e(h["chip"])}</span></p>' if h.get('chip')
               else (f'<p class="eyebrow">{e(h["eyebrow"])}</p>' if h.get('eyebrow') else ''))
        lede = f'<p class="lede">{mdi(h["lede"])}</p>' if h.get('lede') else ''
        return f'<section class="hero"><div class="wrap">{top}<h1>{e(h.get("title", ""))}</h1>{lede}</div></section>\n'

    def page(self, tab):
        return self.hero(tab) + ''.join(self.section(s, tab) for s in tab.get('sections', []))

    def render(self):
        c = self.cfg
        F = json.load(open(ASSETS / 'fonts.json'))
        fontcss = ''.join("@font-face{font-family:'%s';font-style:%s;font-weight:%d;font-display:swap;src:url(data:font/woff2;base64,%s) format('woff2')}\n"
                          % (fam, x['st'], x['w'], x['b']) for fam, xs in F.items() for x in xs)
        css = (ASSETS / 'hub.css').read_text()
        js = (ASSETS / 'hub.js').read_text()
        logo = (ASSETS / 'ahead-logo-reverse.datauri.txt').read_text().strip()
        for k in ('client_logo', 'theme'):
            if c.get(k):
                print(f'note: "{k}" is ignored. Every deliverable hub has the same look; the client is named in the title, cover and footer.')
        pages = {t['key']: self.page(t) for t in self.tabs}
        title = c.get('html_title') or f"{c.get('client')} × AHEAD · {self.tokens['package']}"
        footer = c.get('footer') or f"{c.get('client')} × AHEAD · {self.tokens['package']}{(' · ' + c['date']) if c.get('date') else ''}"
        guide = ('<!--\n  ' + title + '\n  How to add your content:\n'
                 '  1. Search this file for "SLOT:". Each slot names its tab, section and owner.\n'
                 '     Tab keys: ' + ', '.join(t['key'] for t in self.tabs) + '\n'
                 '  2. Replace the whole <div class="ph">...</div> inside your slot with your content\n'
                 '     (paragraphs, lists, tables, callouts, figures). Keep the SLOT and END SLOT comments.\n'
                 '  3. For tables, add one <tr> per row and delete the placeholder row (class="ph-row").\n'
                 '  4. Open the file with ?check at the end of its address to list the placeholders still to fill.\n'
                 '  Look and feel: fixed by the skill. Do not edit the styles or add inline styling; use the building blocks.\n'
                 '  House style: ' + c.get('house_style', 'US English, the client name exactly as the client writes it, no em dashes, no delivery-team names in the content, everything inside this file (no links out).') + '\n-->\n')
        stamp = f'{SKILL_NAME} {VERSION}; look={look_hash()}'
        tpl = c.get('template') or {}
        if tpl.get('name'):
            stamp += f"; template={tpl['name']} {tpl.get('version', '')}".rstrip() + (f" ({tpl['status']})" if tpl.get('status') else '')
        html = ('<!doctype html><html lang="' + e(c.get('lang', 'en-US')) + '"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
                f'<meta name="generator" content="{e(stamp)}">'
                f'<title>{e(title)}</title>\n' + guide + '<style>' + fontcss + css + '</style></head><body id="top">\n'
                f'<header class="top"><div class="brand"><img alt="AHEAD" src="{logo}"></div><nav class="tabs" id="tabs" role="tablist" aria-label="Hub sections">'
                + ''.join(f'<button class="tab" type="button" role="tab" data-tab="{t["key"]}"><span class="tl">{e(t["label"])}</span></button>' for t in self.tabs)
                + '</nav></header>\n<aside class="rail" id="rail" aria-label="On this page"></aside>\n<div id="checkbar" hidden></div>\n<main id="main">\n'
                + ''.join(f'\n<!-- ######################## TAB: {t["key"]} ######################## -->\n<div class="panel" data-tab="{t["key"]}">\n{pages[t["key"]]}</div>\n' for t in self.tabs)
                + f'<footer><div class="wrap">{e(footer)}</div></footer>\n</main>\n<script>{js}</script></body></html>')
        return html


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('config', nargs='?', help='JSON config file')
    ap.add_argument('--preset', help='base | discovery (phase-1) | target-design (phase-2), or a path to a preset JSON')
    ap.add_argument('--out', help='output HTML path')
    ap.add_argument('--guide', help='also write the contributor guide (Markdown) to this path')
    ap.add_argument('--dump-config', help='write the chosen preset/config to this path and stop')
    ap.add_argument('--client'); ap.add_argument('--engagement'); ap.add_argument('--date')
    ap.add_argument('--phase-number', type=int); ap.add_argument('--phase'); ap.add_argument('--next-phase')
    ap.add_argument('--owners', help='"Role=Name;Role=Name" mapping applied on top of the config')
    a = ap.parse_args()
    if a.config:
        path = pathlib.Path(a.config)
    elif a.preset:
        alias = {'base': 'blank', 'general': 'blank', 'blank': 'blank',
                 'discovery': 'discovery', 'phase-1': 'discovery', 'phase1': 'discovery', 'p1': 'discovery',
                 'target-design': 'target-design', 'phase-2': 'target-design', 'phase2': 'target-design', 'p2': 'target-design',
                 'target-platform-design': 'target-design'}
        name = alias.get(a.preset.lower().replace(' ', '-'), a.preset)
        path = pathlib.Path(a.preset) if a.preset.endswith('.json') else ASSETS / 'presets' / f'{name}.json'
        if not path.exists():
            ap.error(f'unknown preset "{a.preset}". Use base, discovery (phase-1) or target-design (phase-2).')
    else:
        ap.error('give a config file or --preset')
    cfg = json.loads(path.read_text(encoding='utf-8'))
    for f in ('client', 'engagement', 'date', 'phase', 'next_phase'):
        if getattr(a, f):
            cfg[f] = getattr(a, f)
    if a.phase_number:
        cfg['phase_number'] = a.phase_number
    if a.owners:
        cfg.setdefault('owners', {}).update(dict(x.split('=', 1) for x in a.owners.split(';') if '=' in x))
    if a.dump_config:
        pathlib.Path(a.dump_config).write_text(json.dumps(cfg, indent=1, ensure_ascii=False), encoding='utf-8')
        print('config written to', a.dump_config)
        return
    if not a.out:
        ap.error('--out is required')
    tpl = cfg.get('template') or {}
    if tpl.get('status') == 'agreed baseline':
        print(f"note: the {tpl['name']} template is version {tpl.get('version')}, an agreed baseline from a phase still in progress. "
              "It may change when that phase closes; tell the user, and see references/changelog.md.")
    hub = Hub(cfg, path.parent)
    html = hub.render()
    pathlib.Path(a.out).write_text(html, encoding='utf-8')
    visible = re.sub(r'<!--[\s\S]*?-->', '', html)
    n_ph = visible.count('<div class="ph" ') + visible.count('<span class="skel"')
    n_slots = html.count('<div class="slot"')
    print(f'written {a.out}  tabs={len(hub.tabs)} slots={n_slots} placeholders={n_ph} size={len(html) // 1024}KB')
    if a.guide:
        from guide import make_guide
        pathlib.Path(a.guide).write_text(make_guide(html), encoding='utf-8')
        print('guide written to', a.guide)


if __name__ == '__main__':
    main()
